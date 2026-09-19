"""
Output Writer
Writes the reconciled data to Excel.

Strategy:
- Main output workbook: Uses win32com (COM automation) to write data into a copy
  of the original master tracker. This preserves ALL Excel internals (Pivot Tables,
  Charts, Macros, Data Models, Conditional Formatting) that openpyxl would silently
  corrupt upon save.
- Reconciliation report: Uses openpyxl to create a simple NEW workbook (no pivots
  or charts to corrupt).
"""

import gc
import os
import shutil
from datetime import datetime
from typing import Any, Dict, List, Optional

import numpy as np
import pandas as pd


class OutputWriteError(Exception):
    """Exception raised for errors during output writing."""
    pass


# ============================================================
# PATH & ARCHIVE HELPERS
# ============================================================

def create_output_path(
    output_dir: str,
    prefix: str = "Master_Updated",
    source_path: Optional[str] = None,
) -> str:
    """
    Create a timestamped output file path inside a Year/Month folder structure.

    Args:
        output_dir: Base output directory.
        prefix: Filename prefix.
        source_path: If provided, the output file extension is copied from this
                     file (e.g. .xlsm stays .xlsm). Otherwise defaults to .xlsx.
    """
    now = datetime.now()
    year_folder = now.strftime("%Y")
    month_folder = now.strftime("%m_%B")

    target_dir = os.path.join(output_dir, year_folder, month_folder)
    os.makedirs(target_dir, exist_ok=True)

    timestamp = now.strftime("%Y%m%d_%H%M%S")

    # Preserve source extension (.xlsm, .xlsx, etc.)
    ext = ".xlsx"
    if source_path:
        _, src_ext = os.path.splitext(source_path)
        if src_ext:
            ext = src_ext

    return os.path.join(target_dir, f"{prefix}_{timestamp}{ext}")


def create_archive_copy(source_path: str, archive_dir: str) -> str:
    """Create a backup copy of the original master tracker."""
    if not source_path or not os.path.isfile(source_path):
        return ""

    now = datetime.now()
    year_folder = now.strftime("%Y")
    month_folder = now.strftime("%m_%B")

    target_dir = os.path.join(archive_dir, year_folder, month_folder)
    os.makedirs(target_dir, exist_ok=True)

    timestamp = now.strftime("%Y%m%d_%H%M%S")
    filename = os.path.basename(source_path)
    name, ext = os.path.splitext(filename)

    archive_path = os.path.join(target_dir, f"{name}_before_{timestamp}{ext}")
    shutil.copy2(source_path, archive_path)
    return archive_path


# ============================================================
# COM TYPE CONVERSION HELPER
# ============================================================

def _to_com_safe(val: Any) -> Any:
    """
    Convert a Python / numpy / pandas value to a type that the Windows COM
    bridge can marshal into an Excel cell without error.

    Handles: numpy scalars, NaN, NaT, pd.Timestamp, None, empty strings.
    """
    if val is None:
        return ""

    # numpy integer types (int8, int16, int32, int64, …)
    if isinstance(val, np.integer):
        return int(val)

    # numpy float types — also catches np.nan
    if isinstance(val, np.floating):
        if np.isnan(val):
            return ""
        return float(val)

    # numpy bool
    if isinstance(val, np.bool_):
        return bool(val)

    # pandas Timestamp or native datetime
    if isinstance(val, (pd.Timestamp, datetime)):
        if pd.isna(val):
            return ""
        # Convert to Excel Serial Date float (days since 1899-12-30)
        # This completely guarantees COM never mangles it into a text string
        delta = pd.Timestamp(val) - pd.Timestamp("1899-12-30")
        return float(delta.total_seconds() / 86400.0)

    # plain Python float NaN
    if isinstance(val, float) and val != val:
        return ""

    # plain Python string — empty / "nan"
    if isinstance(val, str):
        if val.strip() == "" or val.strip().lower() == "nan":
            return ""
        return val

    return val


# ============================================================
# MAIN OUTPUT WORKBOOK (COM-based, pivot-safe)
# ============================================================

def write_output_workbook(
    output_df: pd.DataFrame,
    output_path: str,
    source_workbook_path: Optional[str] = None,
    target_sheets: Optional[List[str]] = None,
    formula_columns: Optional[Dict[str, str]] = None,
) -> str:
    """
    Write the output DataFrame into an Excel workbook using native COM.

    The approach:
      1. Copy the pristine master tracker to output_path (preserving ALL internals).
      2. Open the copy with an invisible Excel instance via COM.
      3. For each target sheet: read headers, detect formulas in row 2,
         clear old data rows, write new data as a 2-D array (fast).
      4. Resize any ListObjects (Excel Tables).
      5. Expand PivotCache source ranges to cover new rows.
      6. RefreshAll → CalculateUntilAsyncQueriesDone → Save → Close.

    This guarantees Pivot Tables, Charts, Macros, and Data Models are never
    destroyed because the file is never touched by openpyxl.
    """
    if not source_workbook_path or not os.path.isfile(source_workbook_path):
        raise OutputWriteError(
            "A source workbook template is required. "
            f"Path: {source_workbook_path}"
        )

    target_sheets = target_sheets or ["Consolidated"]

    # ------------------------------------------------------------------
    # Pre-process DataFrame for COM transfer
    # ------------------------------------------------------------------
    df_write = output_df.copy()

    # Force convert known date columns to real DateTime objects (fixes left-aligned text issue)
    date_columns = ["Created", "Resolved", "Actual Incident Resolve"]
    for col in df_write.columns:
        if col in date_columns or pd.api.types.is_datetime64_any_dtype(df_write[col]):
            # Coerce forces invalid dates to NaT, which we then turn into empty strings
            df_write[col] = pd.to_datetime(df_write[col], errors='coerce')
            df_write[col] = df_write[col].apply(
                lambda x: "" if pd.isna(x) else x.to_pydatetime()
            )

    df_write = df_write.fillna("")

    # ------------------------------------------------------------------
    # 1. Copy pristine template → output path
    # ------------------------------------------------------------------
    try:
        shutil.copy2(source_workbook_path, output_path)
    except Exception as e:
        raise OutputWriteError(f"Cannot copy template to output path: {e}")

    abs_output = os.path.abspath(output_path)

    # ------------------------------------------------------------------
    # 2. Open with COM, write data, refresh pivots
    # ------------------------------------------------------------------
    try:
        import win32com.client
        import pythoncom
    except ImportError:
        raise OutputWriteError(
            "pywin32 is required for pivot-safe writes. "
            "Install with:  pip install pywin32"
        )

    pythoncom.CoInitialize()
    excel = None

    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.DisplayAlerts = False
        excel.Visible = False
        excel.AskToUpdateLinks = False  # prevent dialogs
        excel.AutomationSecurity = 1    # msoAutomationSecurityLow: Bypasses yellow security banner to allow Refresh!

        wb = excel.Workbooks.Open(abs_output, UpdateLinks=0)

        for sheet_name in target_sheets:
            # --- open sheet (skip if missing) ---
            try:
                ws = wb.Worksheets(sheet_name)
            except Exception:
                print(f"  Warning: Sheet '{sheet_name}' not found in "
                      f"workbook, skipping.")
                continue

            # --- read headers from row 1 ---
            last_col = ws.Cells(1, ws.Columns.Count).End(-4159).Column  # xlToLeft
            if last_col < 1:
                continue

            headers: List[str] = []
            for c in range(1, last_col + 1):
                val = ws.Cells(1, c).Value
                headers.append(str(val).strip() if val else "")

            # --- detect formulas from row 2 (skip Duration) ---
            existing_formulas: Dict[int, str] = {}
            old_last_row = ws.Cells(ws.Rows.Count, 1).End(-4162).Row  # xlUp
            if old_last_row >= 2:
                for c in range(1, last_col + 1):
                    cell = ws.Cells(2, c)
                    if cell.HasFormula:
                        h = headers[c - 1]
                        # Duration is computed by Python — do not preserve formula
                        if h.lower() != "duration":
                            existing_formulas[c] = cell.FormulaR1C1

            # --- override with explicit formula config (if any) ---
            if formula_columns:
                for c_idx, h in enumerate(headers, start=1):
                    if h in formula_columns and h.lower() != "duration":
                        existing_formulas[c_idx] = formula_columns[h]

            # --- clear old data rows (row 2 → end) ---
            if old_last_row >= 2:
                ws.Range(
                    ws.Cells(2, 1),
                    ws.Cells(old_last_row, last_col),
                ).ClearContents()

            # --- build 2-D array of new data ---
            num_records = len(df_write)
            if num_records == 0:
                continue

            write_data: List[List[Any]] = []
            for _, row in df_write.iterrows():
                row_data: List[Any] = []
                for c_idx, h in enumerate(headers, start=1):
                    if c_idx in existing_formulas:
                        # Leave formula columns blank for now, applied in second pass
                        row_data.append("")
                    elif h in df_write.columns:
                        row_data.append(_to_com_safe(row[h]))
                    else:
                        row_data.append("")
                write_data.append(row_data)

            # --- write entire array at once (Value preserves real Excel Dates) ---
            target_range = ws.Range(
                ws.Cells(2, 1),
                ws.Cells(1 + num_records, last_col),
            )
            target_range.Value = write_data
            
            # --- Second Pass: Apply preserved R1C1 formulas ---
            for c_idx, f_str in existing_formulas.items():
                f_range = ws.Range(
                    ws.Cells(2, c_idx),
                    ws.Cells(1 + num_records, c_idx)
                )
                f_range.FormulaR1C1 = f_str

            # --- resize ListObjects (Excel Tables) ---
            try:
                for i in range(1, ws.ListObjects.Count + 1):
                    obj = ws.ListObjects(i)
                    new_range = ws.Range(
                        ws.Cells(1, 1),
                        ws.Cells(1 + num_records, last_col),
                    )
                    obj.Resize(new_range)
            except Exception:
                pass  # no tables, or resize not applicable

            print(f"  Sheet '{sheet_name}': {num_records} records written.")

        # --- BULLETPROOF PIVOT UPDATE (HANDLES TABLES AND RANGES) ---
        try:
            # 1. Expand standard ranges for any pivot NOT using a Table
            sheet_ranges = {}
            for sheet_name in target_sheets:
                try:
                    p_ws = wb.Worksheets(sheet_name)
                    last_cell = p_ws.Cells.Find(What="*", SearchOrder=1, SearchDirection=2)
                    if last_cell:
                        last_row = last_cell.Row
                        last_col_cell = p_ws.Cells.Find(What="*", SearchOrder=2, SearchDirection=2)
                        last_col = last_col_cell.Column if last_col_cell else 1
                        sheet_ranges[sheet_name] = f"'{sheet_name}'!R1C1:R{last_row}C{last_col}"
                except Exception:
                    pass

            # 2. Process every single Pivot Cache
            for pc in wb.PivotCaches():
                # Force synchronous refresh to prevent race conditions
                try:
                    pc.BackgroundQuery = False
                except Exception:
                    pass

                # If this cache uses a standard range, expand it
                try:
                    current_src = pc.SourceData
                    if isinstance(current_src, str) and "!" in current_src:
                        src_sheet = current_src.split("!")[0].replace("'", "")
                        if src_sheet in sheet_ranges:
                            pc.SourceData = sheet_ranges[src_sheet]
                except Exception:
                    pass
                
                # Explicitly force this cache to refresh (CRITICAL for 'Table1')
                try:
                    pc.Refresh()
                except Exception as e:
                    print(f"  Warning: Cache refresh failed: {e}")

            # 3. Update all Pivot Tables and force new items to be visible
            for ws in wb.Worksheets:
                for pt in ws.PivotTables():
                    try:
                        # Tell Excel to automatically check the box for new Months/items
                        for pf in pt.PivotFields():
                            try:
                                pf.IncludeNewItemsInFilter = True
                            except Exception:
                                pass
                        pt.Update()
                    except Exception:
                        pass
        except Exception as e:
            print(f"  Warning: Global pivot update failed: {e}")

        # --- refresh all other data connections & formulas ---
        wb.RefreshAll()
        excel.CalculateUntilAsyncQueriesDone()

        wb.Save()
        wb.Close(SaveChanges=True)

        print("  Pivot tables refreshed via COM.")

    except ImportError:
        raise OutputWriteError(
            "pywin32 is required. Install with:  pip install pywin32"
        )
    except OutputWriteError:
        raise
    except Exception as e:
        raise OutputWriteError(f"COM output write failed: {e}")
    finally:
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass
            del excel
        gc.collect()
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass

    return output_path


# ============================================================
# RECONCILIATION REPORT (openpyxl — safe for NEW files)
# ============================================================

def write_reconciliation_report_excel(
    report_data: Dict[str, Any],
    output_path: str,
) -> str:
    """
    Write a reconciliation report to a brand-new Excel workbook.

    Uses openpyxl which is perfectly safe for creating new files (the corruption
    bug only affects files that already contain Pivot Tables / Charts).

    Sheets created:
        - Summary       — key metrics
        - New Incidents — list of newly added incident numbers
        - Changes       — field-level change log
        - IC Lookup Misses — IC values not found in lookup
    """
    from openpyxl import Workbook

    wb = Workbook()

    # --- Summary sheet ---
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.cell(row=1, column=1, value="Metric")
    ws_summary.cell(row=1, column=2, value="Value")
    for i, (key, value) in enumerate(
        report_data.get("summary", {}).items(), 2
    ):
        ws_summary.cell(row=i, column=1, value=str(key))
        ws_summary.cell(row=i, column=2, value=str(value))

    # --- New Incidents sheet ---
    new_incidents = report_data.get("new_incidents", [])
    if new_incidents:
        ws_new = wb.create_sheet("New Incidents")
        ws_new.cell(row=1, column=1, value="Incident Number")
        for i, num in enumerate(new_incidents, 2):
            ws_new.cell(row=i, column=1, value=num)

    # --- Changes sheet ---
    changes = report_data.get("changes", {})
    if changes:
        ws_changes = wb.create_sheet("Changes")
        ws_changes.cell(row=1, column=1, value="Incident Number")
        ws_changes.cell(row=1, column=2, value="Change Details")
        row_idx = 2
        for number, change_list in changes.items():
            for change in change_list:
                ws_changes.cell(row=row_idx, column=1, value=str(number))
                ws_changes.cell(row=row_idx, column=2, value=str(change))
                row_idx += 1

    # --- IC Lookup Misses sheet ---
    ic_misses = report_data.get("ic_misses", [])
    if ic_misses:
        ws_ic = wb.create_sheet("IC Lookup Misses")
        ws_ic.cell(row=1, column=1, value="Details")
        for i, miss in enumerate(ic_misses, 2):
            ws_ic.cell(row=i, column=1, value=str(miss))

    wb.save(output_path)
    return output_path
