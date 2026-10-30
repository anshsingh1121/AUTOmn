"""Excel template loading and audit-result writing helpers."""
from __future__ import annotations
import os
import sys
import openpyxl
from openpyxl.styles import PatternFill
import constants

KWNF_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")  # Light yellow
MM_FILL = PatternFill(fill_type="solid", fgColor="F4CCCC")    # Light red
DM_FILL = PatternFill(fill_type="solid", fgColor="DDEBF7")    # Light blue
MA_FILL = PatternFill(fill_type="solid", fgColor="E4DFEC")    # Light purple

GENERAL_DATA_ROWS = {
    10: "cr_number", 11: "test_date", 12: "change_type", 13: "category",
    14: "closure_status", 15: "related_inc", 16: "related_prb", 17: "jira_issue",
}
QUESTION_ROWS = {
    18: "q1_testing_completed", 19: "q1_testing_completed_new",
    20: "q2_nc_exp_sdlc_ritm", 21: "q3_nc_exp_business_project_patching",
    22: "q4_reason_vs_sdlc", 23: "q5_inc_assoc", 24: "q6_prb_assoc",
    25: "q7_std_template", 26: "q8_differentiation", 27: "q9_approver_check",
    28: "q10_emergency_approval", 29: "q11_fcb_expedited", 30: "q12_normal_low",
    31: "q13_normal_high", 32: "q14_std_approval", 33: "q15_approval_timing",
    34: "q16_major_incident", 35: "q17_post_validation",
    36: "q17_post_validation_new", 37: "q18_issues_doc", 38: "q19_outside_window",
}
STD_Q_ROWS = {
    10: "cr_number", 11: "test_date", 12: "template_ver", 13: "std_proposal",
    14: "category", 15: "closure_status", 16: "related_inc", 17: "jira_issue",
    18: "is_retired", 19: "pre_test_plan", 20: "deployment_plan",
    21: "backout_plan", 22: "verification_plan", 23: "sample_changes",
}
BASIC_DATA_ROWS = {
    3: "test_month", 4: "sample_month", 5: "tester", 6: "population", 7: "samples_tested",
}

def load_template(template_path):
    if not os.path.exists(template_path):
        # Support PyInstaller --onefile bundled template fallback
        if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"):
            bundled = os.path.join(sys._MEIPASS, os.path.basename(template_path))
            if os.path.exists(bundled):
                template_path = bundled
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template file not found at: {template_path}")
    print(f"Loading the template from {template_path}")
    return openpyxl.load_workbook(template_path)

def write_rows(ws, row_mapping, cr_data, col_idx):
    for row, key in row_mapping.items():
        val = cr_data.get(key)
        if val is None or val == "":
            continue
        cell = ws.cell(row=row, column=col_idx)
        cell.value = val
        if val == constants.KWNF:
            cell.fill = KWNF_FILL
        elif val == constants.MM:
            cell.fill = MM_FILL
        elif val == constants.DM:
            cell.fill = DM_FILL
        elif val == constants.MA:
            cell.fill = MA_FILL

def populate_basic_data(workbook, sheet_name, basic_data):
    if sheet_name not in workbook.sheetnames:
        print(f"Warning: Sheet '{sheet_name}' not found in workbook. Skipping.")
        return
    ws = workbook[sheet_name]
    row_data = basic_data[0] if isinstance(basic_data, list) and basic_data else basic_data
    for row, key in BASIC_DATA_ROWS.items():
        value = row_data.get(key)
        if value is not None and value != "":
            ws.cell(row=row, column=2).value = value

def populate_cr_column(ws, cr_data, col_idx, sheet_name):
    if sheet_name == "SC Template Result Sheet":
        write_rows(ws, STD_Q_ROWS, cr_data, col_idx)
    else:
        write_rows(ws, GENERAL_DATA_ROWS, cr_data, col_idx)
        write_rows(ws, QUESTION_ROWS, cr_data, col_idx)

def populate_sheet(workbook, sheet_name, crs_data):
    if sheet_name not in workbook.sheetnames:
        print(f"Warning: Sheet '{sheet_name}' not found in workbook. Skipping.")
        return
    ws = workbook[sheet_name]
    print(f"Populating sheet: {sheet_name} with {len(crs_data)} CRs")
    start_col_idx = 3
    for i, cr_data in enumerate(crs_data):
        populate_cr_column(ws, cr_data, start_col_idx + i, sheet_name)

def populate_sheet_single(workbook, sheet_name, crs_data, start_col_idx=3):
    if sheet_name not in workbook.sheetnames:
        print(f"Warning: Sheet '{sheet_name}' not found in workbook. Skipping.")
        return
    ws = workbook[sheet_name]
    records = crs_data if isinstance(crs_data, list) else [crs_data]
    for i, cr_data in enumerate(records):
        populate_cr_column(ws, cr_data, start_col_idx + i, sheet_name)

def save_output(workbook, output_path):
    os.makedirs("report", exist_ok=True)
    path = output_path if os.path.isabs(output_path) else os.path.join("report", output_path)
    print(f"Saving output to: {path}")
    workbook.save(path)
    return path
