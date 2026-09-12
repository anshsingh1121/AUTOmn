"""
Synthetic Test Data Generator
Creates fictional test Excel files for development and testing.

⚠️ ALL DATA IS SYNTHETIC — NOT CORPORATE DATA ⚠️
"""

import os
import sys
from datetime import datetime, timedelta

# Add project root to path
project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if project_root not in sys.path:
    sys.path.insert(0, project_root)

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def create_master_test_data():
    """Create synthetic Master Tracker test data."""
    wb = Workbook()
    ws = wb.active
    ws.title = "FCB"

    # Headers — matching the documented Master Tracker schema
    headers = [
        "Number", "Short description", "Primary CI/Application",
        "Priority", "State", "Resolved", "Created", "IC",
        "Assignee", "Assignment group", "Caused by Change",
        "Vendor Caused?", "Region", "Bank", "Resolution Time"
    ]
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)

    # Synthetic data rows
    # These cover: active incidents, historical resolved, various priorities/regions
    data = [
        # Row 2: Active incident — will be updated by ServiceNow
        ["INC0010001", "Test app login failure", "App-Portal-001",
         "1 - Critical", "In Progress", None,
         datetime(2026, 3, 15), "Alice Johnson", "Bob Smith",
         "Tier 3 Application Support FCB", "CHG0001001", "", "US", "", None],

        # Row 3: Active incident — will be updated (state change)
        ["INC0010002", "Database timeout on reporting", "DB-Report-002",
         "2 - High", "In Progress", None,
         datetime(2026, 4, 10), "Carol Davis", "Dan Wilson",
         "Tier 2 Database Support", "CHG0001002", "", "India", "", None],

        # Row 4: Historical RESOLVED — NOT in ServiceNow, must be RETAINED
        ["INC0010003", "Network latency spike Jan", "Net-Core-003",
         "3 - Moderate", "Resolved",
         datetime(2026, 1, 20), datetime(2026, 1, 18),
         "Eve Martinez", "Frank Brown",
         "Tier 1 Network Support", "", "", "US", "", None],

        # Row 5: Historical CLOSED — NOT in ServiceNow, must be RETAINED
        ["INC0010004", "Storage alert Q1", "Store-SAN-004",
         "2 - High", "Closed",
         datetime(2026, 2, 5), datetime(2026, 2, 1),
         "Grace Lee", "Hank Taylor",
         "Tier 2 Storage Support", "CHG0001003", "Yes", "India", "", None],

        # Row 6: Active P3 incident — will be updated, tests IC Rule 2
        ["INC0010005", "Minor UI alignment issue", "App-Web-005",
         "3 - Moderate", "Open", None,
         datetime(2026, 5, 22), "Ivy Chen", "Jack White",
         "Tier 1 Application Support FCB", "", "", "US", "", None],

        # Row 7: SVB incident — will be updated, tests Bank rule
        ["INC0010006", "SVB payment processing delay", "Pay-SVB-006",
         "2 - High", "In Progress", None,
         datetime(2026, 6, 1), "Karen Clark", "Leo Harris",
         "Tier 3 Application Support SVB", "", "", "US", "SVB", None],

        # Row 8: Historical P1 RESOLVED — must be RETAINED
        ["INC0010007", "Major outage Feb", "App-Core-007",
         "1 - Critical", "Resolved",
         datetime(2026, 2, 15), datetime(2026, 2, 14),
         "Mike Adams", "Nancy King",
         "Tier 3 Application Support FCB", "CHG0001004", "", "India", "", None],

        # Row 9: Active incident — will be updated
        ["INC0010008", "Email notification failure", "Mail-Srv-008",
         "3 - Moderate", "Open", None,
         datetime(2026, 7, 3), "Oscar Rivera", "Pat Quinn",
         "Tier 2 Application Support FCB", "", "", "US", "", None],
    ]

    for row_idx, row_data in enumerate(data, 2):
        for col_idx, value in enumerate(row_data, 1):
            if col_idx == 15 and row_data[5] is not None and row_data[6] is not None:
                # Resolution Time formula: =[@Resolved]-[@Created]
                ws.cell(row=row_idx, column=col_idx, value="=[@Resolved]-[@Created]")
            elif value is not None:
                ws.cell(row=row_idx, column=col_idx, value=value)

    # Create IC Lookup sheet in same workbook (for reference)
    ws_lookup = wb.create_sheet("IC Lookup")
    lookup_headers = ["IC", "Region"]
    for col, header in enumerate(lookup_headers, 1):
        ws_lookup.cell(row=1, column=col, value=header)

    lookup_data = [
        ["Alice Johnson", "US"],
        ["Bob Smith", "US"],
        ["Carol Davis", "India"],
        ["Dan Wilson", "India"],
        ["Eve Martinez", "US"],
        ["Frank Brown", "US"],
        ["Grace Lee", "India"],
        ["Hank Taylor", "India"],
        ["Ivy Chen", "US"],
        ["Jack White", "US"],
        ["Karen Clark", "US"],
        ["Leo Harris", "US"],
        ["Mike Adams", "India"],
        ["Nancy King", "India"],
        ["Oscar Rivera", "US"],
    ]
    for row_idx, row_data in enumerate(lookup_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws_lookup.cell(row=row_idx, column=col_idx, value=value)

    return wb


def create_servicenow_test_data():
    """Create synthetic ServiceNow export test data."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    headers = [
        "Number", "Short Description", "Configuration Item",
        "Priority", "State", "IC", "Proposed By",
        "Assignment Group", "Assigned To", "Caused by Change",
        "Actual Incident Start", "Actual Incident Resolve", "Vendor"
    ]
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)

    data = [
        # Match INC0010001 — update State to Resolved
        ["INC0010001", "Test app login failure - updated", "App-Portal-001",
         "1 - Critical", "Resolved", "Alice Johnson", "Wendy Thomas",
         "Tier 3 Application Support FCB", "Bob Smith", "CHG0001001",
         datetime(2026, 3, 15), datetime(2026, 9, 1), ""],

        # Match INC0010002 — update Assignment Group, State
        ["INC0010002", "Database timeout on reporting", "DB-Report-002",
         "2 - High", "Resolved", "Carol Davis", "Xavier Young",
         "Tier 3 Database Support", "Dan Wilson", "CHG0001002",
         datetime(2026, 4, 10), datetime(2026, 8, 28), ""],

        # Match INC0010005 — P3, IC blank, tests IC Rule 2 (Proposed By)
        ["INC0010005", "Minor UI alignment issue", "App-Web-005",
         "3 - Moderate", "In Progress", "", "Sam Parker",
         "Tier 1 Application Support FCB", "Jack White", "",
         datetime(2026, 5, 22), None, ""],

        # Match INC0010006 — SVB incident update
        ["INC0010006", "SVB payment processing delay - resolved", "Pay-SVB-006",
         "2 - High", "Resolved", "Karen Clark", "Tom Baker",
         "Tier 3 Application Support SVB", "Leo Harris", "",
         datetime(2026, 6, 1), datetime(2026, 9, 5), "VendorCorp"],

        # Match INC0010008 — IC populated (tests IC Rule 1)
        ["INC0010008", "Email notification failure", "Mail-Srv-008",
         "3 - Moderate", "Resolved", "Oscar Rivera", "Pat Quinn",
         "Tier 2 Application Support FCB", "Pat Quinn", "",
         datetime(2026, 7, 3), datetime(2026, 9, 8), ""],

        # NEW incident — P1, IC populated
        ["INC0010009", "New critical API failure", "API-Gateway-009",
         "1 - Critical", "In Progress", "Alice Johnson", "Uma Patel",
         "Tier 3 Application Support FCB", "Bob Smith", "CHG0001005",
         datetime(2026, 9, 10), None, ""],

        # NEW incident — P3, IC blank (tests IC Rule 2 for new record)
        ["INC0010010", "New minor display glitch", "App-Web-010",
         "3 - Moderate", "Open", "", "Victor Stone",
         "Tier 1 Application Support FCB", "Jack White", "",
         datetime(2026, 9, 11), None, ""],

        # NEW incident — P2, IC blank (tests IC Rule 3)
        ["INC0010011", "New high severity batch failure", "Batch-Proc-011",
         "2 - High", "Open", "", "Wendy Thomas",
         "Tier 2 Application Support FCB", "Dan Wilson", "",
         datetime(2026, 9, 12), None, ""],

        # NEW incident — SVB assignment group (tests Bank rule for new)
        ["INC0010012", "New SVB account sync issue", "Acct-SVB-012",
         "3 - Moderate", "Open", "", "Xavier Young",
         "SVB Go Support", "Leo Harris", "",
         datetime(2026, 9, 12), None, ""],

        # NEW incident — unknown IC (tests IC lookup miss)
        ["INC0010013", "New monitoring alert", "Mon-Dash-013",
         "3 - Moderate", "Open", "Unknown Person", "",
         "Tier 1 Monitoring Support", "Zara Ahmed", "",
         datetime(2026, 9, 12), None, ""],
    ]

    # NOTE: INC0010003, INC0010004, INC0010007 are NOT in ServiceNow
    # They are historical and must be RETAINED in the master

    for row_idx, row_data in enumerate(data, 2):
        for col_idx, value in enumerate(row_data, 1):
            if value is not None and value != "":
                ws.cell(row=row_idx, column=col_idx, value=value)

    return wb


def create_ic_lookup_test_data():
    """Create synthetic IC Lookup test data."""
    wb = Workbook()
    ws = wb.active
    ws.title = "IC Lookup"

    headers = ["IC", "Region"]
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)

    data = [
        ["Alice Johnson", "US"],
        ["Bob Smith", "US"],
        ["Carol Davis", "India"],
        ["Dan Wilson", "India"],
        ["Eve Martinez", "US"],
        ["Frank Brown", "US"],
        ["Grace Lee", "India"],
        ["Hank Taylor", "India"],
        ["Ivy Chen", "US"],
        ["Jack White", "US"],
        ["Karen Clark", "US"],
        ["Leo Harris", "US"],
        ["Mike Adams", "India"],
        ["Nancy King", "India"],
        ["Oscar Rivera", "US"],
        ["Pat Quinn", "US"],
        ["Sam Parker", "US"],
        ["Victor Stone", "US"],
        # "Unknown Person" is intentionally NOT in the lookup
        # to test IC lookup miss reporting
    ]

    for row_idx, row_data in enumerate(data, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    return wb


def main():
    """Generate all synthetic test data files."""
    test_data_dir = os.path.join(project_root, "test_data")
    os.makedirs(test_data_dir, exist_ok=True)

    print("Generating synthetic test data...")
    print("[WARNING] ALL DATA IS SYNTHETIC -- NOT CORPORATE DATA")
    print()

    # Master Tracker
    master_wb = create_master_test_data()
    master_path = os.path.join(test_data_dir, "master_test.xlsx")
    master_wb.save(master_path)
    print(f"  [OK] Master Tracker: {master_path}")

    # ServiceNow Export
    sn_wb = create_servicenow_test_data()
    sn_path = os.path.join(test_data_dir, "servicenow_test.xlsx")
    sn_wb.save(sn_path)
    print(f"  [OK] ServiceNow Export: {sn_path}")

    # IC Lookup
    ic_wb = create_ic_lookup_test_data()
    ic_path = os.path.join(test_data_dir, "ic_lookup_test.xlsx")
    ic_wb.save(ic_path)
    print(f"  [OK] IC Lookup: {ic_path}")

    print()
    print("Test data generation complete.")
    print(f"Files are in: {test_data_dir}")


if __name__ == "__main__":
    main()
